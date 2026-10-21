#!/usr/bin/env python3
"""Verifica os ficheiros .xlsx descarregados da API real pelo teste de integração.

    python3 tests/integracao/verificar_excel.py tests/resultados

Precisa de: pip install openpyxl. Sai com código 1 se alguma verificação falhar.
Os números esperados são os dos dados de exemplo (db/init/04_demo_eventos.sql, jogo 2).
"""
import sys
from pathlib import Path

import openpyxl

pasta = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/resultados")
falhas = []
total = 0


def ok(condicao, mensagem, detalhe=""):
    global total
    total += 1
    print(("OK    " if condicao else "FALHA ") + mensagem + ("" if condicao or not detalhe else f"  -> {detalhe}"))
    if not condicao:
        falhas.append(f"{mensagem} ({detalhe})" if detalhe else mensagem)


def linhas(livro, folha):
    """Linhas da folha sem as completamente vazias."""
    return [l for l in livro[folha].iter_rows(values_only=True) if any(c is not None for c in l)]


def tabela(livro, folha):
    """Folha com cabeçalho na 1.ª linha -> lista de dicionários."""
    todas = linhas(livro, folha)
    cab = todas[0]
    return [dict(zip(cab, l)) for l in todas[1:]]


def abrir(nome):
    caminho = pasta / nome
    ok(caminho.exists(), f"{nome} existe", str(caminho))
    if not caminho.exists():
        return None
    try:
        return openpyxl.load_workbook(caminho)  # falha se o ZIP/XML estiver mal formado
    except Exception as erro:  # noqa: BLE001
        ok(False, f"{nome} abre-se com o openpyxl", repr(erro))
        return None


FOLHAS = ["Resumo", "Jogadores", "Por posição", "Por corredor", "Por distância", "Por parte", "Sistemas defensivos",
          "Guarda-redes", "Defesas na baliza", "Golos na baliza", "Livro dos 7 metros", "Outros eventos", "Eventos"]

ZONAS = ["Alto à esquerda", "Alto ao centro", "Alto à direita", "Meio à esquerda", "Meio ao centro", "Meio à direita",
         "Baixo à esquerda", "Baixo ao centro", "Baixo à direita"]

# ----------------------------------------------------------------------------- jogo completo
jogo = abrir("jogo-2.xlsx")
if jogo:
    ok(jogo.sheetnames == FOLHAS, "jogo-2.xlsx: folhas na ordem esperada", str(jogo.sheetnames))

    resumo = linhas(jogo, "Resumo")
    cabecalho = {l[0]: l[1] for l in resumo if l[0] and l[1] is not None and not isinstance(l[1], (int, float))}
    ok(resumo[0][0] == "Estatísticas do jogo", "Resumo: título", str(resumo[0]))
    ok("Águias do Norte" in str(resumo[1][0]) and "Leões do Sul" in str(resumo[1][0]), "Resumo: nomes das equipas", str(resumo[1]))
    ok(cabecalho.get("Resultado") == "6 – 9", "Resumo: resultado 6 – 9", str(cabecalho.get("Resultado")))
    ok(cabecalho.get("Âmbito") == "Jogo completo", "Resumo: âmbito = Jogo completo", str(cabecalho.get("Âmbito")))

    equipas = {l[0]: l for l in resumo if l[0] in ("Águias do Norte", "Leões do Sul") and isinstance(l[2], (int, float))}
    # colunas: Equipa, Casa/Fora, Remates, Golos, Defesas, Fora, Poste, Bloqueados, Eficácia, ...
    aguias, leoes = equipas.get("Águias do Norte"), equipas.get("Leões do Sul")
    ok(aguias is not None and aguias[2] == 17 and aguias[3] == 6, "Resumo: Águias 17 remates e 6 golos", str(aguias))
    ok(leoes is not None and leoes[2] == 16 and leoes[3] == 9, "Resumo: Leões 16 remates e 9 golos", str(leoes))
    ok(aguias is not None and abs(aguias[8] - 35.3) < 0.06, "Resumo: eficácia das Águias 35,3 %", str(aguias and aguias[8]))
    ok(leoes is not None and abs(leoes[8] - 56.3) < 0.06, "Resumo: eficácia dos Leões 56,3 %", str(leoes and leoes[8]))

    jogadores = tabela(jogo, "Jogadores")
    ok(len(jogadores) == 16, "Jogadores: 16 linhas (8 + 8)", str(len(jogadores)))
    ok(all(j["Convocado"] == "Sim" for j in jogadores), "Jogadores: todos convocados")
    ok(sum(j["Golos"] for j in jogadores) == 15, "Jogadores: soma dos golos = 15", str(sum(j["Golos"] for j in jogadores)))
    ok(sum(j["Remates"] for j in jogadores) == 33, "Jogadores: soma dos remates = 33", str(sum(j["Remates"] for j in jogadores)))

    corredores = [l for l in tabela(jogo, "Por corredor") if l["Equipa"] == "Leões do Sul"]
    ok(sum(l["Remates"] for l in corredores) == 16 and sum(l["Golos"] for l in corredores) == 9,
       "Por corredor: Leões somam 16 remates e 9 golos", str(corredores))
    pontas = [l for l in corredores if "Ponta" in str(l["Corredor"])]
    ok(sum(l["Remates"] for l in pontas) == 5 and sum(l["Golos"] for l in pontas) == 3,
       "Por corredor: pontas dos Leões 3/5", str(pontas))

    partes = tabela(jogo, "Por parte")
    ok(sum(l["Remates"] for l in partes) == 33, "Por parte: 33 remates no total", str(partes))
    distancias = tabela(jogo, "Por distância")
    ok(sum(l["Remates"] for l in distancias) == 33, "Por distância: 33 remates no total", str(distancias))
    posicoes = tabela(jogo, "Por posição")
    ok(sum(l["Remates"] for l in posicoes) == 33, "Por posição: 33 remates no total", str(posicoes))

    grs = tabela(jogo, "Guarda-redes")
    ok(len(grs) == 4, "Guarda-redes: 4 linhas (2 guarda-redes por equipa tiveram remates)", str(len(grs)))
    ok(all(g["Defesas"] + g["Golos sofridos"] == g["Remates à baliza"] for g in grs),
       "Guarda-redes: defesas + golos sofridos = remates à baliza em cada linha", str(grs))
    ok(sum(g["Golos sofridos"] for g in grs) == 15, "Guarda-redes: 15 golos sofridos no total (os 15 golos do jogo)", str(grs))

    # Defesas na baliza: uma linha por guarda-redes, uma coluna por zona (os dados de exemplo têm 11 defesas, todas com local)
    baliza = [l for l in tabela(jogo, "Defesas na baliza") if l.get("N.º") is not None]
    ok(list(tabela(jogo, "Defesas na baliza")[0].keys())[:3 + 9 + 2] == ["Equipa", "N.º", "Nome", *ZONAS, "Sem local", "Total de defesas"],
       "Defesas na baliza: cabeçalho com as 9 zonas", str(list(tabela(jogo, "Defesas na baliza")[0].keys())))
    ok(len(baliza) == 4, "Defesas na baliza: 4 linhas (os 4 guarda-redes que defenderam)", str(len(baliza)))
    ok(sum(b["Total de defesas"] for b in baliza) == 11, "Defesas na baliza: 11 defesas no total", str(baliza))
    ok(all(sum(b[z] for z in ZONAS) + b["Sem local"] == b["Total de defesas"] for b in baliza),
       "Defesas na baliza: zonas + sem local = total em cada guarda-redes", str(baliza))
    ok(all(b["Sem local"] == 0 for b in baliza), "Defesas na baliza: nos dados de exemplo todas as defesas têm local", str(baliza))
    por_zona = [sum(b[z] for b in baliza) for z in ZONAS]
    ok(por_zona == [1, 1, 1, 1, 3, 1, 1, 1, 1], "Defesas na baliza: totais por zona dos dados de exemplo", str(por_zona))
    defesas_gr = {(g["Equipa"], g["Nome"]): g["Defesas"] for g in grs}
    defesas_baliza = {(b["Equipa"], b["Nome"]): b["Total de defesas"] for b in baliza}
    ok(defesas_gr == defesas_baliza, "Defesas na baliza: coincide com as defesas da folha Guarda-redes",
       f"{defesas_gr} != {defesas_baliza}")

    # Golos na baliza: a mesma forma, por guarda-redes que sofreu o golo (os 15 golos de exemplo têm local)
    golos = [l for l in tabela(jogo, "Golos na baliza") if l.get("N.º") is not None]
    ok(list(tabela(jogo, "Golos na baliza")[0].keys())[:3 + 9 + 2] == ["Equipa", "N.º", "Nome", *ZONAS, "Sem local", "Total de golos"],
       "Golos na baliza: cabeçalho com as 9 zonas", str(list(tabela(jogo, "Golos na baliza")[0].keys())))
    ok(sum(b["Total de golos"] for b in golos) == 15, "Golos na baliza: 15 golos no total", str(golos))
    ok(all(sum(b[z] for z in ZONAS) + b["Sem local"] == b["Total de golos"] for b in golos),
       "Golos na baliza: zonas + sem local = total em cada guarda-redes", str(golos))
    ok(all(b["Sem local"] == 0 for b in golos), "Golos na baliza: nos dados de exemplo todos os golos têm local", str(golos))
    golos_por_zona = [sum(b[z] for b in golos) for z in ZONAS]
    ok(golos_por_zona == [2, 2, 1, 1, 2, 2, 2, 1, 2], "Golos na baliza: totais por zona dos dados de exemplo", str(golos_por_zona))
    golos_gr = {(g["Equipa"], g["Nome"]): g["Golos sofridos"] for g in grs}
    golos_baliza = {(b["Equipa"], b["Nome"]): b["Total de golos"] for b in golos}
    ok(golos_gr == golos_baliza, "Golos na baliza: coincide com os golos sofridos da folha Guarda-redes", f"{golos_gr} != {golos_baliza}")

    # Sistemas defensivos: os remates posicionais de exemplo trazem a defesa do adversário (6:0, 5:1 e alguns sem ela)
    sistemas = [l for l in tabela(jogo, "Sistemas defensivos") if l.get("Remates") is not None]
    ok(list(tabela(jogo, "Sistemas defensivos")[0].keys()) == ["Equipa", "Defesa do adversário", "Remates", "Golos", "Defesas", "Fora", "Poste", "Bloqueados", "Eficácia (%)"],
       "Sistemas defensivos: cabeçalho esperado", str(list(tabela(jogo, "Sistemas defensivos")[0].keys())))
    linhas_sis = {(l["Equipa"], l["Defesa do adversário"]): (l["Remates"], l["Golos"]) for l in sistemas}
    ok(linhas_sis == {("Águias do Norte", "6:0"): (9, 3), ("Águias do Norte", "5:1"): (4, 2), ("Águias do Norte", "Sem defesa marcada"): (2, 0),
                      ("Leões do Sul", "6:0"): (9, 5), ("Leões do Sul", "5:1"): (4, 1), ("Leões do Sul", "Sem defesa marcada"): (2, 2)},
       "Sistemas defensivos: remates e golos por equipa e defesa", str(linhas_sis))
    ok(sum(l["Remates"] for l in sistemas) == 30, "Sistemas defensivos: só os 30 remates posicionais (sem contra-ataque nem 7 metros)", str(sum(l["Remates"] for l in sistemas)))
    ok([l["Defesa do adversário"] for l in sistemas if l["Equipa"] == "Águias do Norte"] == ["6:0", "5:1", "Sem defesa marcada"],
       "Sistemas defensivos: pela ordem das defesas, com as sem defesa marcada no fim")
    ok(next(l for l in sistemas if l["Equipa"] == "Leões do Sul" and l["Defesa do adversário"] == "6:0")["Eficácia (%)"] == 55.6, "Sistemas defensivos: eficácia dos Leões contra a 6:0 = 55,6%")

    # Livro dos 7 metros: o jogo de exemplo só tem um 7 metros (golo das Águias, com local na baliza)
    livro = [l for l in tabela(jogo, "Livro dos 7 metros") if l.get("N.º") is not None]
    cab_livro = list(tabela(jogo, "Livro dos 7 metros")[0].keys())
    ok(cab_livro[:10] == ["Equipa", "N.º", "Nome", "Lançamentos", "Golos", "Defendidos", "Fora", "Poste", "Bloqueados", "Eficácia (%)"]
       and cab_livro[10:19] == ZONAS and cab_livro[19:] == ["Sem local", "Zona mais usada"],
       "Livro dos 7 metros: cabeçalho com as 9 zonas", str(cab_livro))
    ok(len(livro) == 1 and livro[0]["Equipa"] == "Águias do Norte" and livro[0]["Lançamentos"] == 1 and livro[0]["Golos"] == 1 and livro[0]["Eficácia (%)"] == 100,
       "Livro dos 7 metros: um lançador, 1 lançamento, 1 golo", str(livro))
    ok(sum(livro[0][z] for z in ZONAS) + livro[0]["Sem local"] == livro[0]["Golos"] + livro[0]["Defendidos"],
       "Livro dos 7 metros: zonas + sem local = golos + defendidos", str(livro))
    ok(isinstance(livro[0]["Zona mais usada"], str) and livro[0]["Zona mais usada"].endswith("(1)"), "Livro dos 7 metros: a zona mais usada vem com a contagem", str(livro[0]["Zona mais usada"]))

    outros = {(l["Equipa"], l["Evento"]): l["Total"] for l in tabela(jogo, "Outros eventos")}
    ok(sum(outros.values()) == 6, "Outros eventos: 6 no total (2 faltas, 2 perdas, 1 dois minutos, 1 timeout)", str(outros))

    eventos = tabela(jogo, "Eventos")
    ok(len(eventos) == 39, "Eventos: 39 linhas (33 remates + 6 outros)", str(len(eventos)))
    ok([e["N.º"] for e in eventos] == list(range(1, 40)), "Eventos: numeração 1..39")
    ok(sum(1 for e in eventos if e["Evento"] == "Remate") == 33, "Eventos: 33 remates")
    ok(sum(1 for e in eventos if e["Resultado"] == "Golo") == 15, "Eventos: 15 golos")
    defesas = [e for e in eventos if e["Resultado"] == "Defesa"]
    ok(len(defesas) == 11 and all(e["Zona da baliza"] in ZONAS for e in defesas),
       "Eventos: as 11 defesas têm a zona da baliza", str([e["Zona da baliza"] for e in defesas]))
    ok(all(isinstance(e["Baliza X (%)"], (int, float)) and 0 <= e["Baliza X (%)"] <= 100
           and isinstance(e["Baliza Y (%)"], (int, float)) and 0 <= e["Baliza Y (%)"] <= 100 for e in defesas),
       "Eventos: coordenadas da baliza entre 0 e 100 nas defesas")
    golos_ev = [e for e in eventos if e["Resultado"] == "Golo"]
    ok(len(golos_ev) == 15 and all(e["Zona da baliza"] in ZONAS for e in golos_ev),
       "Eventos: os 15 golos têm a zona da baliza", str([e["Zona da baliza"] for e in golos_ev]))
    ok(all(e["Zona da baliza"] is None and e["Baliza X (%)"] is None for e in eventos if e["Resultado"] not in ("Defesa", "Golo")),
       "Eventos: só as defesas e os golos têm local na baliza")
    posicionais = [e for e in eventos if e["Situação"] == "Posicional" and e["Resultado"] is not None]
    ok(sum(1 for e in posicionais if e["Defesa do adversário"] is not None) == 26 and {e["Defesa do adversário"] for e in posicionais} == {"6:0", "5:1", None},
       "Eventos: a coluna «Defesa do adversário» traz a defesa dos 26 remates posicionais que a têm", str({e["Defesa do adversário"] for e in posicionais}))
    ok(all(e["Defesa do adversário"] is None for e in eventos if e["Situação"] != "Posicional" or e["Resultado"] is None),
       "Eventos: só os remates posicionais têm a defesa do adversário")
    ok(all(isinstance(e["Guarda-redes"], str) and e["Guarda-redes"].startswith("#") for e in defesas),
       "Eventos: as defesas indicam o guarda-redes (#nº nome)", str([e["Guarda-redes"] for e in defesas][:3]))
    tempos = [e["Tempo"] for e in eventos]
    ok(all(isinstance(t, str) and len(t) >= 5 and t[2] == ":" for t in tempos), "Eventos: tempo no formato mm:ss", str(tempos[:3]))

    # Nenhuma célula de texto pode ser interpretada como fórmula
    perigosas = [c.coordinate for ws in jogo for linha in ws.iter_rows() for c in linha
                 if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("="))]
    ok(not perigosas, "nenhuma célula é uma fórmula", str(perigosas[:5]))

# ----------------------------------------------------------------------------- só uma equipa
equipa = abrir("jogo-2-equipa-2.xlsx")
if equipa:
    ok(equipa.sheetnames == FOLHAS, "jogo-2-equipa-2.xlsx: mesmas folhas", str(equipa.sheetnames))
    resumo = {l[0]: l[1] for l in linhas(equipa, "Resumo") if l[0] and isinstance(l[1], str)}
    ok("Leões do Sul" in str(resumo.get("Âmbito")), "equipa: âmbito menciona os Leões do Sul", str(resumo.get("Âmbito")))
    jogadores = tabela(equipa, "Jogadores")
    ok(len(jogadores) == 8 and {j["Equipa"] for j in jogadores} == {"Leões do Sul"}, "equipa: 8 jogadores, todos dos Leões do Sul", str(len(jogadores)))
    eventos = tabela(equipa, "Eventos")
    ok(len(eventos) == 19 and {e["Equipa"] for e in eventos} == {"Leões do Sul"}, "equipa: 19 eventos, todos dos Leões do Sul", str(len(eventos)))
    corredores = tabela(equipa, "Por corredor")
    ok({l["Equipa"] for l in corredores} == {"Leões do Sul"} and sum(l["Remates"] for l in corredores) == 16,
       "equipa: por corredor só dos Leões, 16 remates", str(corredores))
    grs = tabela(equipa, "Guarda-redes")
    ok({g["Equipa"] for g in grs} == {"Leões do Sul"}, "equipa: só o guarda-redes dos Leões", str(grs))
    baliza = [l for l in tabela(equipa, "Defesas na baliza") if l.get("N.º") is not None]
    ok(len(baliza) >= 1 and {b["Equipa"] for b in baliza} == {"Leões do Sul"}, "equipa: defesas na baliza só dos guarda-redes dos Leões", str(baliza))
    ok(sum(b["Total de defesas"] for b in baliza) == sum(g["Defesas"] for g in grs), "equipa: defesas na baliza = defesas da folha Guarda-redes", str(baliza))
    sistemas_eq = [l for l in tabela(equipa, "Sistemas defensivos") if l.get("Remates") is not None]
    ok({l["Equipa"] for l in sistemas_eq} == {"Leões do Sul"} and sum(l["Remates"] for l in sistemas_eq) == 15,
       "equipa: sistemas defensivos só dos Leões (15 remates posicionais)", str(sistemas_eq))
    ok([l for l in tabela(equipa, "Livro dos 7 metros") if l.get("N.º") is not None] == [], "equipa: os Leões não têm 7 metros no livro")
    golos_baliza = [l for l in tabela(equipa, "Golos na baliza") if l.get("N.º") is not None]
    ok({b["Equipa"] for b in golos_baliza} == {"Leões do Sul"}, "equipa: golos na baliza só dos guarda-redes dos Leões", str(golos_baliza))
    ok(sum(b["Total de golos"] for b in golos_baliza) == sum(g["Golos sofridos"] for g in grs), "equipa: golos na baliza = golos sofridos da folha Guarda-redes", str(golos_baliza))

# ----------------------------------------------------------------------------- jogo novo, quase vazio
novo = abrir("jogo-novo.xlsx")
if novo:
    ok(novo.sheetnames == FOLHAS, "jogo-novo.xlsx: mesmas folhas", str(novo.sheetnames))
    eventos = tabela(novo, "Eventos")
    ok(len(eventos) >= 1, "jogo novo: tem os eventos registados no teste", str(len(eventos)))
    # o teste de integração registou 2 defesas ao "GR B": uma em baixo à direita (83,5 / 91,2) e outra sem local
    defesas = [e for e in eventos if e["Resultado"] == "Defesa"]
    ok(len(defesas) == 2, "jogo novo: 2 defesas registadas", str(len(defesas)))
    com_local = [e for e in defesas if e["Zona da baliza"] is not None]
    ok(len(com_local) == 1 and com_local[0]["Zona da baliza"] == "Baixo à direita"
       and abs(com_local[0]["Baliza X (%)"] - 83.5) < 0.001 and abs(com_local[0]["Baliza Y (%)"] - 91.2) < 0.001,
       "jogo novo: a defesa com local aparece em 'Baixo à direita' (83,5 / 91,2)", str(com_local))
    ok(sum(1 for e in defesas if e["Zona da baliza"] is None and e["Baliza X (%)"] is None) == 1, "jogo novo: a outra defesa fica sem local")
    baliza = [l for l in tabela(novo, "Defesas na baliza") if l.get("N.º") is not None]
    gr_b = [b for b in baliza if b["Nome"] == "GR B"]
    ok(len(gr_b) == 1 and gr_b[0]["Baixo à direita"] == 1 and gr_b[0]["Sem local"] == 1 and gr_b[0]["Total de defesas"] == 2,
       "jogo novo: GR B tem 1 defesa em baixo à direita e 1 sem local", str(gr_b))

# ----------------------------------------------------------------------------- Excel das Estatísticas da equipa
# Cada ficheiro é conferido contra os números que a API dá para os mesmos filtros (guardados pelo teste de integração em
# equipa-estatisticas*.json): as somas são refeitas aqui, de forma independente do código do exportador.
import json

FOLHAS_EQUIPA = ["Resumo", "Jogos", "Ataque", "Guarda-redes", "Bola e disciplina", "Jogo a jogo"]
CONTADORES = ["remates", "golos", "defendidos", "fora", "poste", "bloqueados", "seteMetros", "golosSeteMetros", "contraAtaques",
              "golosContraAtaque", "perdasBola", "faltasOfensivas", "doisMinutos", "vermelhos", "defesas", "golosSofridos",
              "seteMetrosSofridos", "defesasSeteMetros"]


def ler_json(nome):
    caminho = pasta / nome
    ok(caminho.exists(), f"{nome} existe", str(caminho))
    return json.loads(caminho.read_text(encoding="utf-8")) if caminho.exists() else None


def jogos_do_filtro(dados, local="todos", ultimos=None):
    jogos = sorted(dados["jogos"], key=lambda j: (j["dataHora"], j["jogoId"]))
    if local == "casa":
        jogos = [j for j in jogos if j["emCasa"]]
    elif local == "fora":
        jogos = [j for j in jogos if not j["emCasa"]]
    return jogos[-ultimos:] if ultimos else jogos


def verificar_equipa(ficheiro, dados, local="todos", ultimos=None):
    livro = abrir(ficheiro)
    if not livro or not dados:
        return
    jogos = jogos_do_filtro(dados, local, ultimos)
    ids = {j["jogoId"] for j in jogos}
    linhas_filtradas = [l for l in dados["linhas"] if l["jogoId"] in ids]
    plantel = {p["id"]: p for p in dados["plantel"]}
    totais = {}
    for l in linhas_filtradas:
        t = totais.setdefault(l["jogadorId"], {c: 0 for c in CONTADORES} | {"jogos": 0})
        t["jogos"] += 1
        for c in CONTADORES:
            t[c] += l[c]
    soma = lambda c: sum(t[c] for t in totais.values())  # noqa: E731
    nome = f"{ficheiro}:"

    ok(livro.sheetnames == FOLHAS_EQUIPA, f"{nome} folhas na ordem esperada", str(livro.sheetnames))

    resumo = linhas(livro, "Resumo")
    cabecalho = {l[0]: l[1] for l in resumo[:-2] if l[0] and l[1] is not None and not isinstance(l[1], (int, float))}   # sem a tabela do fim
    ok(str(resumo[0][0]).startswith("Estatísticas de "), f"{nome} Resumo: título com o nome da equipa", str(resumo[0]))
    esperado_jogos = {"todos": "Todos", "casa": "Em casa", "fora": "Fora"}[local]
    ok(cabecalho.get("Jogos") == esperado_jogos, f"{nome} Resumo: «Jogos» = {esperado_jogos}", str(cabecalho.get("Jogos")))
    ok(cabecalho.get("Período") == (f"Últimos {ultimos} jogos" if ultimos else "Toda a época"), f"{nome} Resumo: período", str(cabecalho.get("Período")))
    resultados = ["V" if j["golos"] > j["golosAdversario"] else "D" if j["golos"] < j["golosAdversario"] else "E" for j in jogos]
    marcados, sofridos = sum(j["golos"] for j in jogos), sum(j["golosAdversario"] for j in jogos)
    esperado = [len(jogos), resultados.count("V"), resultados.count("E"), resultados.count("D"), marcados, sofridos, marcados - sofridos]
    ok(list(resumo[-1][:7]) == esperado, f"{nome} Resumo: jogos, V-E-D, golos marcados e sofridos e diferença = {esperado}", str(resumo[-1]))

    folha_jogos = tabela(livro, "Jogos")
    ok(len(folha_jogos) == len(jogos), f"{nome} Jogos: {len(jogos)} jogos", str(len(folha_jogos)))
    ok([l["Adversário"] for l in folha_jogos] == [j["adversarioNome"] for j in jogos], f"{nome} Jogos: os adversários, do mais antigo para o mais recente")
    ok([l["Golos"] for l in folha_jogos] == [j["golos"] for j in jogos] and [l["Golos sofridos"] for l in folha_jogos] == [j["golosAdversario"] for j in jogos],
       f"{nome} Jogos: os golos de cada jogo")
    ok([l["Resultado"] for l in folha_jogos] == resultados and all(l["Casa/Fora"] == ("Casa" if j["emCasa"] else "Fora") for l, j in zip(folha_jogos, jogos)),
       f"{nome} Jogos: o resultado e o local de cada jogo")
    ok(sum(l["Remates"] for l in folha_jogos) == soma("remates") and sum(l["Perdas de bola"] for l in folha_jogos) == soma("perdasBola"),
       f"{nome} Jogos: os remates e as perdas de bola somam o que somam os jogadores")

    ataque = [l for l in tabela(livro, "Ataque") if l["Jogador"] != "Equipa"]
    equipa = next((l for l in tabela(livro, "Ataque") if l["Jogador"] == "Equipa"), None)
    por_nome = {p["nome"]: p for p in plantel.values()}
    certos = all(l["Jogador"] in por_nome and totais.get(por_nome[l["Jogador"]]["id"], {}).get("golos") == l["Golos"]
                 and totais[por_nome[l["Jogador"]]["id"]]["remates"] == l["Remates"] and totais[por_nome[l["Jogador"]]["id"]]["jogos"] == l["Jogos"]
                 and l["#"] == por_nome[l["Jogador"]]["numero"] for l in ataque)
    ok(certos, f"{nome} Ataque: jogos, golos e remates de cada jogador são os somados das linhas", str(ataque))
    ok([l["Golos"] for l in ataque] == sorted((l["Golos"] for l in ataque), reverse=True), f"{nome} Ataque: o melhor marcador primeiro")
    ok(all(l["Golos"] <= l["Remates"] for l in ataque) and all((l["Eficácia (%)"] is None) == (l["Remates"] == 0) for l in ataque), f"{nome} Ataque: a eficácia só existe com remates")
    ok(all(l["Remates"] == 0 or abs(l["Eficácia (%)"] - 100 * l["Golos"] / l["Remates"]) < 0.06 for l in ataque), f"{nome} Ataque: eficácia = golos / remates")
    ok(equipa is not None and equipa["Golos"] == soma("golos") and equipa["Remates"] == soma("remates") and equipa["Jogos"] == len(jogos),
       f"{nome} Ataque: a linha da equipa soma os jogadores e conta os jogos da equipa", str(equipa))
    ok(equipa is not None and equipa["Golos"] == marcados, f"{nome} Ataque: os golos da equipa são os golos dos jogos ({marcados})", str(equipa))

    gr = [l for l in tabela(livro, "Guarda-redes") if l["Jogador"] != "Equipa"]
    gr_equipa = next((l for l in tabela(livro, "Guarda-redes") if l["Jogador"] == "Equipa"), None)
    ok(all(l["Defesas"] + l["Golos sofridos"] == l["Remates à baliza"] for l in gr), f"{nome} Guarda-redes: defesas + golos sofridos = remates à baliza")
    ok(all(totais.get(por_nome[l["Jogador"]]["id"], {}).get("defesas") == l["Defesas"] for l in gr), f"{nome} Guarda-redes: as defesas de cada um são as somadas das linhas", str(gr))
    ok(gr_equipa is not None and gr_equipa["Defesas"] == soma("defesas") and gr_equipa["Golos sofridos"] == soma("golosSofridos"), f"{nome} Guarda-redes: linha da equipa", str(gr_equipa))

    disciplina = [l for l in tabela(livro, "Bola e disciplina") if l["Jogador"] != "Equipa"]
    ok(len(disciplina) == len(totais) and sum(l["Perdas de bola"] for l in disciplina) == soma("perdasBola") and sum(l["Vermelhos"] for l in disciplina) == soma("vermelhos"),
       f"{nome} Bola e disciplina: um jogador por linha e as somas certas", str(disciplina))
    ok([l["Exclusões de 2 min"] for l in disciplina] == sorted((l["Exclusões de 2 min"] for l in disciplina), reverse=True), f"{nome} Bola e disciplina: por ordem de exclusões")

    jogo_a_jogo = tabela(livro, "Jogo a jogo")
    ok(len(jogo_a_jogo) == len(linhas_filtradas), f"{nome} Jogo a jogo: uma linha por jogador e jogo ({len(linhas_filtradas)})", str(len(jogo_a_jogo)))
    ok(sum(l["Golos"] for l in jogo_a_jogo) == soma("golos") and sum(l["Defesas"] for l in jogo_a_jogo) == soma("defesas"), f"{nome} Jogo a jogo: os golos e as defesas somam o mesmo")
    ok(livro["Jogo a jogo"].auto_filter.ref is not None and livro["Jogo a jogo"].freeze_panes == "A2", f"{nome} Jogo a jogo: filtros e primeira linha fixa")


equipa_json = ler_json("equipa-estatisticas.json")
parte1_json = ler_json("equipa-estatisticas-parte-1.json")
parte2_json = ler_json("equipa-estatisticas-parte-2.json")
if equipa_json:
    ok(len(equipa_json["jogos"]) >= 1 and len(equipa_json["linhas"]) >= 2, "equipa-estatisticas.json: a equipa tem jogos e linhas para conferir", f"{len(equipa_json['jogos'])} jogos")
    verificar_equipa("equipa-estatisticas.xlsx", equipa_json)
    verificar_equipa("equipa-estatisticas-casa.xlsx", equipa_json, local="casa")
    verificar_equipa("equipa-estatisticas-ultimos-2.xlsx", equipa_json, ultimos=2)
if parte2_json:
    verificar_equipa("equipa-estatisticas-parte-2.xlsx", parte2_json)
if parte1_json:
    verificar_equipa("equipa-estatisticas-fora.xlsx", parte1_json, local="fora", ultimos=3)
    if equipa_json and parte2_json:
        ok(len(parte1_json["jogos"]) <= len(equipa_json["jogos"]) and len(parte2_json["jogos"]) <= len(equipa_json["jogos"]), "as partes têm no máximo os jogos do jogo todo")

print()
if falhas:
    print(f"{len(falhas)} verificação(ões) falharam:")
    for f in falhas:
        print(" -", f)
    resumo = "%0A".join(falhas[:20]).replace("\r", "")
    print(f"::error title=Verificação dos ficheiros Excel::{resumo}")
    sys.exit(1)
print(f"Todas as {total} verificações dos ficheiros Excel passaram.")
